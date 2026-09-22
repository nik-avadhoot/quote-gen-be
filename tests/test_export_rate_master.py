"""/export RATE MASTER contract: grades added in the app reach the workbook.

Runs the real export route (auth unwrapped) against the shipped template.
A grade the template does not list used to be dropped, so every layer using it
exported with a blank material rate and the workbook costed differently.
"""
import io
import pathlib
import sys

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import server  # noqa: E402

PASSES, FAILURES = 0, []


def check(condition, label):
    global PASSES
    if condition:
        PASSES += 1
        print(f"ok   - {label}")
    else:
        FAILURES.append(label)
        print(f"FAIL - {label}")


rates = [
    {"code": "16", "desc": "16 BF Kraft", "price": 31.5, "disc": 1, "freight": 0},
    {"code": "25", "desc": "25 BF Kraft", "price": 41.25, "disc": 1, "freight": 0},
    {"code": "30GY", "desc": "30 BF Golden Yellow", "price": 47, "disc": 1.5, "freight": 0.5},
    {"code": "24GY", "desc": "24 BF GY", "price": 41.5, "disc": 1, "freight": 0, "interest": 2},
]
spec = {"client": "Probe", "delivery": "Nagpur", "plant": "Nagpur", "rowType": "Box",
        "material_code": "P1", "L": 300, "W": 200, "H": 150, "ply": 3, "interest": 1.25,
        "layers": {"TOP": {"code": "30GY", "gsm": 150}}}

with server.app.test_request_context(
        "/export", method="POST", json={"items": [{"spec": spec}], "rates": rates, "freight": {}}):
    response = server.export_xlsx.__wrapped__()
    response.direct_passthrough = False
    wb = openpyxl.load_workbook(io.BytesIO(response.get_data()))

rm, cbb = wb["RATE MASTER"], wb["CBB+PP"]

# The 24GY row's position is a fact about the shipped template, not this test —
# it moved once already (row 16 -> 17) when the Product Owner added a grade by
# hand. Look it up instead of hardcoding it, the same way server.py does.
_row_24gy = next(r for r in range(7, 25) if rm.cell(r, 1).value == "24GY")

check(rm["A34"].value == "30GY" and rm["C34"].value == 47 and rm["E34"].value == 1.5
      and rm["F34"].value == 0.5, "RM-1 a grade missing from the template is appended below the notes")
check(rm["D34"].value == "=C34*$B$4" and rm["G34"].value == "=C34+D34-E34+F34",
      "RM-2 the appended grade carries the template's credit and effective-rate formulas")
check(rm.cell(_row_24gy, 4).value == f"=C{_row_24gy}*0.02",
      "RM-3 a per-grade supplier credit replaces $B$4 for that grade only")
check(rm["D7"].value == "=C7*$B$4", "RM-4 a grade without its own credit keeps the sheet-wide $B$4")
# The GSM surcharge table was relocated from $A$26:$C$29 to $K$7:$M$10 by the
# same hand-edit (CBB+PP's VLOOKUP was updated to match) to make room for the
# grade the append logic must never write into or above. Check where it lives
# now, not its former address.
check([rm.cell(r, 11).value for r in range(7, 11)] == [0, 100, 101, 201] and rm["M8"].value == 1.5,
      "RM-5 the GSM surcharge table (now at $K$7:$M$10) is untouched")
check(abs(cbb["BJ3"].value - 0.0125) < 1e-12, "RM-6 BJ3 carries the item's interest")
# The hand-added grade arrived as the NUMBER 25. CBB+PP writes layer codes as
# text and VLOOKUP(...,0) never matches text against a number, so the row looked
# updated and still resolved to a blank rate.
_row_25 = next(r for r in range(7, 25) if str(rm.cell(r, 1).value).strip() == "25")
check(isinstance(rm.cell(_row_25, 1).value, str) and rm.cell(_row_25, 1).value == "25",
      "RM-7 a matched grade code is written back as text, not left numeric")

print(f"\n{PASSES} passed, {len(FAILURES)} failed")
if FAILURES:
    sys.exit(1)
print("Export RATE MASTER contract PASS")
