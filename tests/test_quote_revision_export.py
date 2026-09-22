"""The governed Quote export: an APPROVED revision becomes the master workbook.

Product Owner, 2026-09-22: the same workbook as the working export, and the only
export that carries a permanent Quote reference.

Drives the real route with the auth decorator unwrapped and the caller reads
stubbed, so the assembly and the guards are exercised rather than described. The
workbook is opened afterwards and its cells are read: this fails if a frozen
value stops reaching the sheet.
"""
import io
import pathlib
import sys

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import server  # noqa: E402

PASSES, FAILURES = 0, []


def check(condition, label):
    global PASSES
    if condition:
        PASSES += 1
        print(f"ok   - {label}")
    else:
        FAILURES.append(label)
        print(f"FAIL - {label}")


SNAPSHOT = {
    "pricing_basis_release_id": 7,
    "freight_set_version_id": 11,
    "effective_inputs": {
        "entered": {
            "length_mm": 675, "width_mm": 450, "height_mm": 282, "ply": 5, "ups": 2,
            "box_type": "RSC", "flute_f1": "B", "flute_f2": "C", "item_name": "DTY 5PLY 36512",
            "layers": {
                "TOP": {"code": "35GY", "gsm": 170}, "F1": {"code": "16", "gsm": 120},
                "L1": {"code": "24", "gsm": 170}, "F2": {"code": "16", "gsm": 120},
                "L2": {"code": "24", "gsm": 170},
            },
            "add_ons": {"printing": 2.5, "stitching": 0, "coating": 0, "handling": 0,
                        "moq_charge": 0, "packing": 0, "other": 0, "unloading": 0},
        },
        "resolved": {
            "waste": {"value": 5}, "conv": {"value": 7}, "margin": {"value": 8},
            "interest": {"value": 1.0}, "freight": {"value": 1.25},
        },
    },
}

QUOTE = {
    "quote_reference": "NAG/QUO/2026-27/00007",
    "batch": {"id": 501, "plant": {"name": "Nagpur"}, "customer_family": {"name": "Indo Rama"}},
    "revisions": [{
        "id": 91, "workflow_status": "approved", "quote_date": "2026-09-22",
        "offer_validity_to": "2026-10-22",
        "approved_by_actor": {"display_name": "Snehal"},
        "items": [{"batch_row_lineage_id": 4001, "calculation_snapshot": SNAPSHOT}],
    }],
}

TABLE_ROWS = {
    "batch_rows": [{"lineage_id": 4001, "material_code": "36512", "row_type": "Box"}],
    "pricing_basis_releases": [{"id": 7, "rate_set_version_id": 21, "freight_set_version_id": 11}],
    "rate_entries": [
        {"grade_code": "35GY", "description": "35 BF Golden Yellow", "price": 52,
         "discount": 1, "freight": 0, "interest_pct": None},
        {"grade_code": "16", "description": "16 BF Kraft", "price": 31.5,
         "discount": 1, "freight": 0, "interest_pct": None},
        {"grade_code": "24", "description": "24 BF Kraft", "price": 40,
         "discount": 1, "freight": 0, "interest_pct": None},
    ],
    "freight_entries": [{"origin_plant_id": 1, "destination_location_id": 55, "rate": 1.25}],
    "plants": [{"id": 1, "name": "Nagpur"}],
    "customer_locations": [{"id": 55, "location_name": "Butibori"}],
}


class _Query:
    """Enough of the PostgREST builder for the reads this route makes."""

    def __init__(self, table):
        self.table = table

    def select(self, *_a, **_k):
        return self

    def eq(self, *_a, **_k):
        return self

    def in_(self, *_a, **_k):
        return self

    def limit(self, *_a, **_k):
        return self


class _Client:
    def table(self, name):
        return _Query(name)


def run(revision_status="approved", quote=None, stub_rows=None):
    """Call the route with auth unwrapped and every caller read stubbed."""
    rows = dict(TABLE_ROWS)
    rows.update(stub_rows or {})
    payload = quote if quote is not None else QUOTE
    payload["revisions"][0]["workflow_status"] = revision_status

    original_reads = (server._read_quote_workspace, server._optional_caller_rows,
                      server.get_supabase_for_caller)
    server._read_quote_workspace = lambda *a, **k: payload
    server._optional_caller_rows = lambda query: (rows.get(query.table, []), False)
    server.get_supabase_for_caller = lambda *_a, **_k: _Client()
    try:
        with server.app.test_request_context(f"/quotes/revisions/91/export"):
            server.g.access_token = "stub"
            response = server.export_quote_revision_route.__wrapped__(91)
    finally:
        (server._read_quote_workspace, server._optional_caller_rows,
         server.get_supabase_for_caller) = original_reads
    return response


# ── an approved revision produces the workbook ─────────────────────────────
response = run()
workbook = None
if hasattr(response, "get_data"):
    response.direct_passthrough = False
    workbook = openpyxl.load_workbook(io.BytesIO(response.get_data()))
check(workbook is not None, "QE-1 an approved revision returns a workbook")

if workbook:
    cbb, rm = workbook["CBB+PP"], workbook["RATE MASTER"]
    check(cbb["D4"].value and "NAG/QUO/2026-27/00007" in str(cbb["D4"].value),
          "QE-2 the permanent Quote reference reaches the sheet")
    check(cbb["D2"].value == "Indo Rama" and cbb["B3"].value == "Nagpur",
          "QE-3 customer and producing plant come from the Batch, not from a typed field")
    check(cbb["C7"].value == "36512" and cbb["B7"].value == "Box",
          "QE-4 the row carries its material code and row type")
    check(cbb["F7"].value == 675 and cbb["G7"].value == 450 and cbb["H7"].value == 282,
          "QE-5 frozen dimensions reach the row")
    check(cbb["AD7"].value == "35GY" and cbb["AE7"].value == 170,
          "QE-6 frozen paper layers reach the row")
    # The whole point of freezing: these are the snapshot's resolved values.
    check(abs(cbb["BJ3"].value - 0.01) < 1e-12, "QE-7 the frozen interest is written, not recomputed")
    check(abs(cbb["AY3"].value - 0.05) < 1e-12 and cbb["BA3"].value == 7,
          "QE-8 frozen waste and conversion are written")
    check(abs(cbb["BM3"].value - 0.08) < 1e-12, "QE-9 the frozen margin is written")
    # Frozen Rate Set, not today's master.
    row_35gy = next((r for r in range(7, 40) if str(rm.cell(r, 1).value).strip() == "35GY"), None)
    check(row_35gy is not None and rm.cell(row_35gy, 3).value == 52,
          "QE-10 paper prices come from the Rate Set frozen with the Quote")

# ── the guards ─────────────────────────────────────────────────────────────
for status in ("draft", "submitted"):
    refused = run(revision_status=status)
    body = refused[0].get_json() if isinstance(refused, tuple) else refused.get_json()
    # The house error shape: INVALID_INPUT is 400 with the text under "error".
    check(isinstance(refused, tuple) and refused[1] == 400
          and body.get("error_code") == "INVALID_INPUT"
          and "approved" in (body.get("error") or ""),
          f"QE-11 a {status} revision is refused — no reference exists yet")

missing_snapshot = {
    **QUOTE,
    "revisions": [{**QUOTE["revisions"][0],
                   "items": [{"batch_row_lineage_id": 4001, "calculation_snapshot": None}]}],
}
refused = run(quote=missing_snapshot)
check(isinstance(refused, tuple) and refused[1] == 400,
      "QE-12 an invisible snapshot refuses the whole document rather than exporting a partial one")

refused = run(stub_rows={"rate_entries": []})
check(isinstance(refused, tuple) and refused[1] == 400,
      "QE-13 an unreadable frozen Rate Set refuses rather than exporting rate-less rows")

print(f"\n{PASSES} passed, {len(FAILURES)} failed")
if FAILURES:
    sys.exit(1)
print("Governed Quote export contract PASS")
